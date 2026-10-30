# Part of Odoo. See LICENSE file for full copyright and licensing details.
import base64
import csv
import io
import re
import time
import zipfile
from pathlib import Path
from unittest.mock import patch

import openpyxl

from odoo.tests import tagged
from odoo.tools import html2plaintext

from .common import LiwanCrmCase
from odoo.addons.liwan_api.controllers.common import ApiError
from odoo.addons.liwan_api.tests.common import set_probe_feature
from odoo.addons.liwan_crm.models import intake as intake_module
from odoo.addons.liwan_crm.models.intake import (
    FILE_BAD_CSV,
    FILE_BAD_TYPE,
    FILE_BAD_XLSX,
    FILE_DUPLICATE_HEADER,
    FILE_EMPTY,
    FILE_MESSAGES,
    FILE_NO_ROWS,
    FILE_NOT_UTF8,
    FILE_TOO_LARGE,
    FILE_TOO_MANY_COLUMNS,
    FILE_TOO_MANY_ROWS,
    FILE_UNPACKED_TOO_LARGE,
    MAX_BYTES,
    MAX_COLUMNS,
    MAX_UNCOMPRESSED_BYTES,
    file_error,
)

PREVIEW_URL = '/api/v1/crm/imports/preview'
COLUMNS_URL = '/api/v1/crm/imports/columns'
IMPORTS_URL = '/api/v1/crm/imports'
MAPPINGS_URL = '/api/v1/crm/import-mappings'
DATA = Path(__file__).parent / 'data'
MAPPING = {'name': 'Name', 'phone': 'Phone', 'email': 'Email', 'notes': 'Notes', 'package': 'Package'}
HEADER = 'Name,Phone\n'


def encoded(content):
    return base64.b64encode(content).decode()


def csv_of_size(size):
    """A valid CSV of exactly ``size`` bytes: customer rows with a long unmapped ``Pad`` column."""
    head = b'Name,Phone,Pad\n'
    line = 100_000
    count, extra = divmod(size - len(head), line)
    first = b'Sara,0501234567,' + b'x' * (line + extra - 17) + b'\n'
    other = b'Sara,0501234567,' + b'x' * (line - 17) + b'\n'
    return head + first + other * (count - 1)


def xlsx_bytes(cells, dimension=None, extra=None):
    """A workbook with ``cells`` ({(row, column): value}); ``dimension`` forges the sheet's declared
    range and ``extra`` is the byte size of a highly compressible junk member."""
    book = openpyxl.Workbook()
    for (row, column), value in cells.items():
        book.active.cell(row=row, column=column, value=value)
    buffer = io.BytesIO()
    book.save(buffer)
    if dimension is None and extra is None:
        return buffer.getvalue()
    out = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if dimension and item.filename == 'xl/worksheets/sheet1.xml':
                data = re.sub(rb'<dimension ref="[^"]*"/>', f'<dimension ref="{dimension}"/>'.encode(), data)
            target.writestr(item.filename, data)
        if extra:
            target.writestr('docProps/junk.bin', b'\0' * extra)
    return out.getvalue()


def csv_bytes(*rows, header=HEADER):
    return (header + ''.join(f'{row}\n' for row in rows)).encode()


@tagged('post_install', '-at_install', 'liwan')
class TestCrmImport(LiwanCrmCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company.country_id = cls.env.ref('base.sa')
        cls.source = cls.env['utm.source'].create({'name': 'Spreadsheet'})
        cls.package = cls.make_package(price=250.0, name='Basic')
        cls.set_settings(distribution_mode='manual')
        cls.intake = cls.env['liwan.crm.intake']

    def body(self, content, file_name='customers.csv', mapping=None, **extra):
        return {
            'file_name': file_name, 'content_base64': encoded(content),
            'mapping': mapping or {'name': 'Name', 'phone': 'Phone'}, 'source_id': self.source.id, **extra,
        }

    def preview(self, content, user=None, **kwargs):
        return self.as_user(user or self.manager).api('POST', PREVIEW_URL, json=self.body(content, **kwargs))

    def commit(self, content, user=None, **kwargs):
        return self.as_user(user or self.manager).api('POST', IMPORTS_URL, json=self.body(content, **kwargs))

    def lead_of(self, phone):
        return self.env['crm.lead'].with_context(active_test=False).search([('phone_sanitized', '=', phone)])

    # Reading files

    def test_csv_and_xlsx_produce_identical_rows(self):
        from_csv = self.intake._liwan_read_rows('customers.csv', (DATA / 'customers.csv').read_bytes())
        from_xlsx = self.intake._liwan_read_rows('customers.xlsx', (DATA / 'customers.xlsx').read_bytes())
        self.assertEqual(from_csv, from_xlsx)
        self.assertEqual(len(from_csv), 3)
        self.assertEqual(set(from_csv[0]), set(MAPPING.values()))
        self.assertEqual(from_csv[1]['Phone'], '507654321')

    def test_xlsx_numbers_become_text_without_a_decimal_point(self):
        book = openpyxl.Workbook()
        book.active.append(['Name', 'Phone'])
        book.active.append(['Sara', 501234567.0])
        book.active.append(['Omar', 507654321])
        buffer = io.BytesIO()
        book.save(buffer)
        rows = self.intake._liwan_read_rows('numbers.xlsx', buffer.getvalue())
        self.assertEqual([row['Phone'] for row in rows], ['501234567', '507654321'])

    def test_byte_order_mark_is_not_part_of_the_first_header(self):
        rows = self.intake._liwan_read_rows('customers.csv', (DATA / 'customers.csv').read_bytes())
        self.assertIn('Name', rows[0])

    def test_semicolon_delimited_csv_is_read(self):
        rows = self.intake._liwan_read_rows('x.csv', b'Name;Phone\nSara;0501234567\n')
        self.assertEqual(rows, [{'Name': 'Sara', 'Phone': '0501234567'}])

    def test_quoted_punctuation_does_not_pick_the_delimiter(self):
        rows = self.intake._liwan_read_rows('x.csv', b'"Customer name, company";Phone\nSara;0501234567\n')
        self.assertEqual(rows, [{'Customer name, company': 'Sara', 'Phone': '0501234567'}])
        rows = self.intake._liwan_read_rows('x.csv', b'"Name, first"\tPhone\nSara\t0501234567\n')
        self.assertEqual(rows, [{'Name, first': 'Sara', 'Phone': '0501234567'}])

    def test_windows_arabic_encoding_is_refused_with_a_clear_error(self):
        content = 'الاسم,الهاتف\nسارة,0501234567\n'.encode('cp1256')
        with self.assertRaises(ApiError) as caught:
            self.intake._liwan_read_rows('customers.csv', content)
        self.assertEqual((caught.exception.status, caught.exception.field), (422, 'file'))
        response = self.preview(content, mapping={'name': 'x', 'phone': 'y'})
        self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(response.json()['field'], 'file')

    def test_blank_rows_are_dropped(self):
        rows = self.intake._liwan_read_rows('x.csv', csv_bytes('Sara,0501234567', ',', '', ' , ', 'Omar,0507654321', ',', ','))
        self.assertEqual([row['Name'] for row in rows], ['Sara', 'Omar'])

    def test_row_limit_is_five_thousand(self):
        full = csv_bytes(*[f'Customer {number},0501234567' for number in range(5000)])
        self.assertEqual(len(self.intake._liwan_read_rows('x.csv', full)), 5000)
        over = csv_bytes(*[f'Customer {number},0501234567' for number in range(5001)])
        with self.assertRaises(ApiError) as caught:
            self.intake._liwan_read_rows('x.csv', over)
        self.assertEqual((caught.exception.status, caught.exception.field), (422, 'file'))

    def test_size_limit_is_five_megabytes(self):
        at_limit = csv_of_size(MAX_BYTES)
        self.assertEqual(len(at_limit), MAX_BYTES)
        self.assertEqual(self.preview(at_limit).status_code, 200)
        # One byte over passes the encoded-length check and is refused after decoding; twice the
        # limit is refused on its encoded length alone.
        for content in (csv_of_size(MAX_BYTES + 1), csv_of_size(2 * MAX_BYTES)):
            with self.subTest(size=len(content)):
                response = self.preview(content)
                self.assertEqual(response.status_code, 422, response.text)
                self.assertEqual(response.json()['field'], 'file')
        with self.assertRaises(ApiError) as caught:
            self.intake._liwan_read_rows('x.csv', csv_of_size(MAX_BYTES + 1))
        self.assertEqual((caught.exception.status, caught.exception.field), (422, 'file'))

    def test_unreadable_files_are_refused(self):
        for name, content in (('customers.xlsx', b'not a workbook'), ('customers.txt', b'Name,Phone\n'), ('customers.csv', b''),
                              ('customers.csv', b'Name,Phone\n')):
            with self.subTest(name=name, content=content[:8]):
                response = self.preview(content, file_name=name)
                self.assertEqual(response.status_code, 422, response.text)
                self.assertIn(response.json()['field'], ('file', 'file_name'))

    def test_bad_base64_is_refused(self):
        body = {**self.body(b'x'), 'content_base64': '***'}
        response = self.as_manager().api('POST', PREVIEW_URL, json=body)
        self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(response.json()['field'], 'content_base64')

    def test_xlsx_with_a_huge_uncompressed_size_is_refused_before_it_is_opened(self):
        small = xlsx_bytes({(1, 1): 'Name', (1, 2): 'Phone', (2, 1): 'Sara', (2, 2): '0501234567'})
        self.assertEqual(len(self.intake._liwan_read_rows('x.xlsx', small)), 1)
        bomb = xlsx_bytes({(1, 1): 'Name', (1, 2): 'Phone', (2, 1): 'Sara', (2, 2): '0501234567'},
                          extra=MAX_UNCOMPRESSED_BYTES)
        self.assertLess(len(bomb), MAX_BYTES)
        with self.assertRaises(ApiError) as caught:
            self.intake._liwan_read_rows('x.xlsx', bomb)
        self.assertEqual((caught.exception.status, caught.exception.field), (422, 'file'))

    def test_forged_wide_dimension_does_not_make_the_reader_scan_every_column(self):
        cells = {(1, 1): 'Name', (1, 2): 'Phone', **{(row, 1): f'Customer {row}' for row in range(2, 4002)},
                 **{(row, 2): '0501234567' for row in range(2, 4002)}}
        forged = xlsx_bytes(cells, dimension='A1:XFD1048576')
        started = time.monotonic()
        rows = self.intake._liwan_read_rows('x.xlsx', forged)
        self.assertLess(time.monotonic() - started, 10)
        self.assertEqual(len(rows), 4000)
        self.assertEqual(set(rows[0]), {'Name', 'Phone'})

    def test_a_row_far_below_the_data_hits_the_scan_limit(self):
        cells = {(1, 1): 'Name', (1, 2): 'Phone', (2, 1): 'Sara', (2, 2): '0501234567', (1_000_000, 1): 'Far'}
        started = time.monotonic()
        response = self.preview(xlsx_bytes(cells), file_name='x.xlsx')
        self.assertLess(time.monotonic() - started, 10)
        self.assertEqual((response.status_code, response.json()['field']), (422, 'file'))

    def test_more_columns_than_the_limit_are_refused(self):
        header = ','.join(f'Column {number}' for number in range(MAX_COLUMNS + 1))
        response = self.preview(f'{header}\n'.encode(), mapping={'name': 'Column 0', 'phone': 'Column 1'})
        self.assertEqual((response.status_code, response.json()['field']), (422, 'file'))
        xlsx = xlsx_bytes({(1, number + 1): f'Column {number}' for number in range(MAX_COLUMNS + 1)} | {(2, 1): 'x'})
        response = self.preview(xlsx, file_name='x.xlsx', mapping={'name': 'Column 0', 'phone': 'Column 1'})
        self.assertEqual((response.status_code, response.json()['field']), (422, 'file'))

    def test_duplicate_headers_are_refused_by_every_route(self):
        content = b'Name,Phone,Name\nSara,0501234567,Omar\n'
        sheet = xlsx_bytes({(1, 1): 'Name', (1, 2): 'Phone', (1, 3): 'Name', (2, 1): 'Sara'})
        for url in (PREVIEW_URL, IMPORTS_URL, COLUMNS_URL):
            for name, data in (('x.csv', content), ('x.xlsx', sheet)):
                with self.subTest(url=url, name=name):
                    body = self.body(data, file_name=name)
                    if url == COLUMNS_URL:
                        body = {key: body[key] for key in ('file_name', 'content_base64')}
                    response = self.as_manager().api('POST', url, json=body)
                    self.assertEqual(response.status_code, 422, response.text)
                    self.assertEqual((response.json()['field'], response.json()['details']), ('file', {'reason': FILE_DUPLICATE_HEADER}))
        self.assertFalse(self.lead_of('+966501234567'))

    def test_every_file_refusal_carries_its_machine_code(self):
        columns = ','.join(f'Column {number}' for number in range(MAX_COLUMNS + 1))
        rows = csv_bytes(*[f'Customer {number},0501234567' for number in range(5001)])
        sheet = {(1, 1): 'Name', (1, 2): 'Phone', (2, 1): 'Sara', (2, 2): '0501234567'}
        bomb = xlsx_bytes(sheet, extra=MAX_UNCOMPRESSED_BYTES)
        cases = (
            (FILE_TOO_LARGE, 'x.csv', csv_of_size(MAX_BYTES + 1), 'file'),
            (FILE_TOO_LARGE, 'x.csv', csv_of_size(2 * MAX_BYTES), 'file'),
            (FILE_TOO_MANY_ROWS, 'x.csv', rows, 'file'),
            (FILE_TOO_MANY_COLUMNS, 'x.csv', f'{columns}\n'.encode(), 'file'),
            (FILE_UNPACKED_TOO_LARGE, 'x.xlsx', bomb, 'file'),
            (FILE_DUPLICATE_HEADER, 'x.csv', b'Name,Phone,Name\nSara,0501234567,Omar\n', 'file'),
            (FILE_NOT_UTF8, 'x.csv', 'الاسم,الهاتف\n'.encode('cp1256'), 'file'),
            (FILE_BAD_XLSX, 'x.xlsx', b'not a workbook', 'file'),
            (FILE_EMPTY, 'x.csv', b'', 'file'),
            (FILE_NO_ROWS, 'x.csv', b'Name,Phone\n', 'file'),
            (FILE_BAD_TYPE, 'x.txt', b'Name,Phone\n', 'file_name'),
        )
        for code, name, content, field in cases:
            for url in (PREVIEW_URL, COLUMNS_URL):
                with self.subTest(code=code, name=name, url=url, size=len(content)):
                    body = self.body(content, file_name=name)
                    if url == COLUMNS_URL:
                        body = {key: body[key] for key in ('file_name', 'content_base64')}
                    response = self.as_manager().api('POST', url, json=body)
                    self.assertEqual(response.status_code, 422, response.text)
                    self.assertEqual((response.json()['field'], response.json()['details']), (field, {'reason': code}))

    def test_a_csv_the_reader_cannot_parse_carries_its_code(self):
        # Odoo lifts the csv field-size limit, so a parser error is simulated instead of crafted.
        with patch('odoo.addons.liwan_crm.models.intake.csv.reader', side_effect=csv.Error('bad')):
            with self.assertRaises(ApiError) as caught:
                self.intake._liwan_read_rows('x.csv', b'Name,Phone\nSara,0501234567\n')
        self.assertEqual((caught.exception.status, caught.exception.field, caught.exception.details), (422, 'file', {'reason': FILE_BAD_CSV}))

    # Row numbers

    def test_row_numbers_are_source_rows_even_after_a_blank_row(self):
        self.make_lead(self.agent2, contact_name='Existing', phone='+966501112222')
        csv_content = csv_bytes('Sara,0501234567', ',', 'Omar,12', 'Huda,0501112222', '', ',', 'Lina,0501234567')
        sheet = {(1, 1): 'Name', (1, 2): 'Phone', (2, 1): 'Sara', (2, 2): '0501234567', (4, 1): 'Omar', (4, 2): '12',
                 (5, 1): 'Huda', (5, 2): '0501112222', (8, 1): 'Lina', (8, 2): '0501234567'}
        for name, content in (('x.csv', csv_content), ('x.xlsx', xlsx_bytes(sheet))):
            with self.subTest(name=name):
                body = self.preview(content, file_name=name).json()
                self.assertEqual(body['invalid_rows'], [{'row': 4, 'reason': 'invalid_phone'}])
                self.assertEqual(body['duplicate_rows'], [5, 8])

    def test_a_csv_cell_with_a_line_break_is_one_row(self):
        content = b'Name,Phone\n"Sara\nAl",0501234567\nOmar,12\n'
        body = self.preview(content).json()
        self.assertEqual(body['invalid_rows'], [{'row': 3, 'reason': 'invalid_phone'}])

    # Columns

    def test_columns_route_lists_headers_and_counts_rows(self):
        for name, content in (('x.csv', (DATA / 'customers.csv').read_bytes()), ('x.xlsx', (DATA / 'customers.xlsx').read_bytes())):
            with self.subTest(name=name):
                response = self.as_manager().api('POST', COLUMNS_URL, json={'file_name': name, 'content_base64': encoded(content)})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json(), {'columns': ['Name', 'Phone', 'Email', 'Notes', 'Package'], 'rows': 3})

    def test_columns_route_has_the_same_gates_and_limits(self):
        body = {'file_name': 'x.csv', 'content_base64': encoded(csv_bytes('Sara,0501234567'))}
        self.assertEqual(self.as_agent().api('POST', COLUMNS_URL, json=body).status_code, 403)
        bad = self.as_manager().api('POST', COLUMNS_URL, json={**body, 'content_base64': encoded('الاسم\n'.encode('cp1256'))})
        self.assertEqual((bad.status_code, bad.json()['field']), (422, 'file'))
        extra = self.as_manager().api('POST', COLUMNS_URL, json={**body, 'source_id': 1})
        self.assertEqual((extra.status_code, extra.json()['field']), (422, 'source_id'))
        set_probe_feature(self, key='crm.import', enabled=False)
        self.assertEqual(self.as_manager().api('POST', COLUMNS_URL, json=body).status_code, 403)

    # Preview

    def test_phone_variants_normalise_and_repeats_in_the_file_are_duplicates(self):
        variants = ['0501234567', '501234567', '+966501234567', '00966501234567', '966501234567', '050-123-4567',
                    '٠٥٠١٢٣٤٥٦٧', '۰۵۰۱۲۳۴۵۶۷']
        content = csv_bytes(*[f'Customer {number},{phone}' for number, phone in enumerate(variants)])
        body = self.preview(content).json()
        self.assertEqual((body['new'], body['duplicate'], body['invalid']), (1, 7, 0))
        self.assertEqual(body['duplicate_rows'], list(range(3, 10)))
        self.commit(content)
        self.assertEqual(len(self.lead_of('+966501234567')), 1)

    def test_a_file_the_export_produced_imports_cleanly(self):
        """The export prefixes an E.164 phone with a tab so spreadsheets keep the plus; it must not survive the import."""
        self.make_lead(self.agent2, contact_name='Exported customer', phone='+966501119999')
        exported = self.as_manager().api('GET', '/api/v1/crm/export')
        self.assertEqual(exported.status_code, 200, exported.text)
        self.assertIn('\t+966501119999', exported.content.decode('utf-8-sig'))
        columns = {'name': 'name', 'phone': 'phone'}
        body = self.preview(exported.content, mapping=columns).json()
        self.assertEqual((body['new'], body['duplicate'], body['invalid']), (0, 1, 0))
        fresh = exported.content.replace(b'+966501119999', b'+966501118888')
        body = self.preview(fresh, mapping=columns).json()
        self.assertEqual((body['new'], body['duplicate'], body['invalid']), (1, 0, 0))
        self.commit(fresh, mapping=columns)
        self.assertEqual(len(self.lead_of('+966501118888')), 1)

    def test_existing_lost_or_archived_opportunity_counts_as_duplicate(self):
        lost = self.make_lead(self.agent2, contact_name='Lost customer', phone='+966501112222')
        lost.write({'active': False})
        body = self.preview(csv_bytes('Sara,0501112222', 'Omar,0507654321')).json()
        self.assertEqual((body['new'], body['duplicate'], body['invalid']), (1, 1, 0))
        self.assertEqual(body['duplicate_rows'], [2])

    def test_preview_discloses_row_numbers_only(self):
        self.make_lead(self.agent2, contact_name='Hidden customer', phone='+966501112222')
        response = self.preview(csv_bytes('Sara,0501112222'))
        self.assertEqual(set(response.json()), {'new', 'duplicate', 'invalid', 'invalid_rows', 'duplicate_rows'})
        self.assertEqual(response.json()['duplicate_rows'], [2])
        self.assertNotIn('Hidden', response.text)
        self.assertNotIn('CRM Agent Two', response.text)

    def test_invalid_rows_carry_stable_reason_codes(self):
        content = csv_bytes(',0501234567', 'Sara,', 'Omar,12', 'Huda,0501234567', 'Lina,0507654321,')
        rows = self.preview(content).json()['invalid_rows']
        self.assertEqual(rows, [
            {'row': 2, 'reason': 'missing_name'}, {'row': 3, 'reason': 'missing_phone'},
            {'row': 4, 'reason': 'invalid_phone'},
        ])

    def test_invalid_email_is_a_reason(self):
        content = b'Name,Phone,Email\nSara,0501234567,not-an-email\n'
        body = self.preview(content, mapping={'name': 'Name', 'phone': 'Phone', 'email': 'Email'}).json()
        self.assertEqual(body['invalid_rows'], [{'row': 2, 'reason': 'invalid_email'}])

    def test_package_matches_by_trimmed_case_insensitive_name(self):
        content = b'Name,Phone,Package\nSara,0501234567, BASIC \nOmar,0507654321,Gold\nHuda,0501112222,\n'
        mapping = {'name': 'Name', 'phone': 'Phone', 'package': 'Package'}
        body = self.preview(content, mapping=mapping).json()
        self.assertEqual((body['new'], body['invalid']), (2, 1))
        self.assertEqual(body['invalid_rows'], [{'row': 3, 'reason': 'unknown_package'}])
        self.commit(content, mapping=mapping)
        sara = self.lead_of('+966501234567')
        self.assertEqual(sara.liwan_package_id, self.package)
        self.assertEqual(sara.liwan_list_price, 250.0)
        self.assertEqual(sara.liwan_final_price, 0.0)
        self.assertFalse(self.lead_of('+966501112222').liwan_package_id)

    def test_archived_and_foreign_packages_are_unknown(self):
        archived = self.make_package(name='Retired')
        archived.active = False
        foreign = self.env['product.template'].create({
            'name': 'Foreign', 'liwan_crm_package': True, 'company_id': self.other_company.id,
        })
        content = f'Name,Phone,Package\nSara,0501234567,{archived.name}\nOmar,0507654321,{foreign.name}\n'.encode()
        body = self.preview(content, mapping={'name': 'Name', 'phone': 'Phone', 'package': 'Package'}).json()
        self.assertEqual(body['invalid'], 2)

    def test_mapped_column_missing_from_the_file_is_refused(self):
        response = self.preview(csv_bytes('Sara,0501234567'), mapping={'name': 'Name', 'phone': 'Mobile'})
        self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(response.json()['field'], 'mapping')

    def test_mapping_requires_name_and_phone_and_known_keys(self):
        for mapping in ({'name': 'Name'}, {'phone': 'Phone'}, {'name': 'Name', 'phone': 'Phone', 'city': 'City'},
                        {'name': '', 'phone': 'Phone'}, 'Name,Phone'):
            with self.subTest(mapping=mapping):
                response = self.preview(csv_bytes('Sara,0501234567'), mapping=mapping)
                self.assertEqual(response.status_code, 422, response.text)
                self.assertEqual(response.json()['field'], 'mapping')

    # Commit

    def test_commit_counts_match_the_preview_and_the_batch_keeps_them(self):
        self.make_lead(self.agent2, contact_name='Existing', phone='+966501112222')
        content = csv_bytes('Sara,0501234567', 'Omar,0507654321', 'Huda,0501112222', 'Lina,0501234567', ',0501119999', 'Mona,12')
        preview = self.preview(content).json()
        self.assertEqual((preview['new'], preview['duplicate'], preview['invalid']), (2, 2, 2))
        response = self.commit(content)
        self.assertEqual(response.status_code, 200, response.text)
        batch = self.env['liwan.crm.import.batch'].browse(response.json()['id'])
        self.assertEqual((batch.created_count, batch.duplicate_count, batch.invalid_count), (2, 2, 2))
        self.assertEqual(batch.company_id, self.company)
        self.assertEqual(batch.user_id, self.manager)
        self.assertEqual(batch.file_name, 'customers.csv')
        self.assertEqual(batch.source_id, self.source)
        self.assertEqual(response.json(), {
            'id': batch.id, 'file_name': 'customers.csv', 'user': {'id': self.manager.id, 'name': self.manager.name},
            'source': {'id': self.source.id, 'name': 'Spreadsheet'}, 'created': 2, 'duplicate': 2, 'invalid': 2,
            'unassigned': 2, 'created_at': response.json()['created_at'],
        })

    def test_commit_creates_opportunities_in_the_new_stage_with_the_batch_and_source(self):
        response = self.commit(b'Name,Phone,Email,Notes\nSara,0501234567,sara@example.com,Call after noon\n',
                               mapping={'name': 'Name', 'phone': 'Phone', 'email': 'Email', 'notes': 'Notes'})
        self.assertEqual(response.status_code, 200, response.text)
        lead = self.lead_of('+966501234567')
        self.assertEqual(lead.type, 'opportunity')
        self.assertEqual(lead.stage_id, self.env.ref('liwan_crm.stage_new'))
        self.assertEqual(lead.source_id, self.source)
        self.assertEqual(lead.contact_name, 'Sara')
        self.assertEqual(lead.email_from, 'sara@example.com')
        self.assertIn('Call after noon', lead.description)
        self.assertEqual(lead.company_id, self.company)
        self.assertEqual(lead.liwan_import_batch_id.id, response.json()['id'])
        self.assertFalse(lead.user_id)

    def test_commit_revalidates_instead_of_trusting_a_preview(self):
        content = csv_bytes('Sara,0501234567')
        self.assertEqual(self.preview(content).json()['new'], 1)
        self.make_lead(self.agent2, contact_name='Taken meanwhile', phone='+966501234567')
        body = self.commit(content).json()
        self.assertEqual((body['created'], body['duplicate']), (0, 1))

    def test_a_phone_taken_during_the_commit_is_a_duplicate_not_an_error(self):
        self.make_lead(self.agent2, contact_name='Racing customer', phone='+966501234567')
        Lead = type(self.env['crm.lead'])
        with patch.object(Lead, '_liwan_find_duplicate', return_value=False):
            response = self.commit(csv_bytes('Sara,0501234567', 'Omar,0507654321'))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual((response.json()['created'], response.json()['duplicate']), (1, 1))
        self.assertEqual(len(self.lead_of('+966507654321')), 1)

    def test_source_is_required_and_must_be_active(self):
        content = csv_bytes('Sara,0501234567')
        body = self.body(content)
        del body['source_id']
        missing = self.as_manager().api('POST', IMPORTS_URL, json=body)
        self.assertEqual((missing.status_code, missing.json()['field']), (422, 'source_id'))
        self.source.active = False
        retired = self.commit(content)
        self.assertEqual((retired.status_code, retired.json()['field']), (422, 'source_id'))
        self.assertFalse(self.lead_of('+966501234567'))

    def test_preview_reflects_the_errors_of_the_commit(self):
        content = csv_bytes('Sara,0501234567')
        body = self.body(content)
        del body['source_id']
        missing = self.as_manager().api('POST', PREVIEW_URL, json=body)
        self.assertEqual((missing.status_code, missing.json()['field']), (422, 'source_id'))
        wrong_mode = self.preview(content, user_id=self.agent.id)
        self.assertEqual((wrong_mode.status_code, wrong_mode.json()['field']), (422, 'user_id'))
        self.set_settings(distribution_mode='on_import')
        no_agent = self.preview(content)
        self.assertEqual((no_agent.status_code, no_agent.json()['field']), (422, 'user_id'))
        self.assertEqual(self.preview(content, user_id=self.agent.id).status_code, 200)

    def test_mapping_and_profile_together_are_refused(self):
        profile = self.env['liwan.crm.import.mapping'].create({
            'company_id': self.company.id, 'name': 'Sheet', 'name_column': 'Name', 'phone_column': 'Phone',
        })
        for url in (PREVIEW_URL, IMPORTS_URL):
            response = self.as_manager().api('POST', url, json=self.body(csv_bytes('Sara,0501234567'), mapping_profile=profile.id))
            self.assertEqual((response.status_code, response.json()['field']), (422, 'mapping_profile'), url)

    def test_agent_is_chosen_only_in_on_import_mode(self):
        response = self.commit(csv_bytes('Sara,0501234567'), user_id=self.agent.id)
        self.assertEqual((response.status_code, response.json()['field']), (422, 'user_id'))

    def test_on_import_assigns_the_chosen_agent_and_rings_the_bell(self):
        self.set_settings(distribution_mode='on_import')
        response = self.commit(csv_bytes('Sara,0501234567', 'Omar,0507654321'), user_id=self.agent.id)
        self.assertEqual(response.status_code, 200, response.text)
        leads = self.lead_of('+966501234567') | self.lead_of('+966507654321')
        self.assertEqual(leads.user_id, self.agent)
        self.assertEqual(response.json()['unassigned'], 0)
        for lead in leads:
            self.assertFalse(self.assignment_mails(lead))
        first = min(leads, key=lambda lead: lead.id)
        bells = self.bell_rows(self.agent, first, 'crm.assigned')
        self.assertEqual(len(bells), 1, 'one bell for the whole batch')
        self.assertEqual(html2plaintext(bells.body).strip(), first.contact_name)
        self.assertEqual(bells.liwan_crm_count, 2)
        for other in leads - first:
            self.assertFalse(self.bell_rows(self.agent, other, 'crm.assigned'))

    def test_on_import_needs_an_eligible_agent(self):
        self.set_settings(distribution_mode='on_import')
        content = csv_bytes('Sara,0501234567')
        self.assertEqual(self.commit(content).json()['field'], 'user_id')
        self.assertEqual(self.commit(content, user_id=self.outsider.id).status_code, 422)
        self.assertFalse(self.lead_of('+966501234567'))

    def test_round_robin_spreads_the_batch(self):
        self.set_settings(distribution_mode='round_robin')
        response = self.commit(csv_bytes('Sara,0501234567', 'Omar,0507654321', 'Huda,0501112222', 'Lina,0501119999'))
        self.assertEqual(response.status_code, 200, response.text)
        leads = self.env['crm.lead'].search([('liwan_import_batch_id', '=', response.json()['id'])])
        self.assertEqual(len(leads), 4)
        self.assertTrue(all(leads.mapped('user_id')))
        self.assertGreater(len(set(leads.mapped('user_id').ids)), 1)

    def test_round_robin_import_works_for_an_importer_without_the_work_capability(self):
        importer = self.custom_user('crm_importer', ['crm.leads.read', 'crm.leads.assign', 'crm.leads.import'])
        self.set_settings(distribution_mode='round_robin')
        response = self.commit(csv_bytes('Sara,0501234567', 'Omar,0507654321'), user=importer)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['unassigned'], 0)
        self.assertTrue(self.env['liwan.crm.rotation'].search([('company_id', '=', self.company.id)]).last_user_id)

    def test_import_starts_in_the_first_open_stage_even_when_the_won_stage_sorts_first(self):
        self.env.ref('liwan_crm.stage_won').sequence = 1
        self.env.ref('liwan_crm.stage_new').active = False
        self.assertEqual(self.commit(csv_bytes('Sara,0501234567')).status_code, 200)
        self.assertEqual(self.lead_of('+966501234567').stage_id, self.env.ref('liwan_crm.stage_contacted'))

    def test_manual_mode_leaves_the_batch_in_the_pool(self):
        response = self.commit(csv_bytes('Sara,0501234567'))
        self.assertEqual(response.json()['unassigned'], 1)

    def test_xlsx_upload_through_the_routes(self):
        content = (DATA / 'customers.xlsx').read_bytes()
        response = self.commit(content, file_name='customers.xlsx', mapping=MAPPING)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['created'], 3)
        self.assertEqual(self.lead_of('+966507654321').contact_name, 'Omar')
        self.assertEqual(self.lead_of('+966501234567').liwan_package_id, self.package)

    def test_an_unknown_file_refusal_code_is_a_programming_error(self):
        with self.assertRaisesRegex(ValueError, 'file_typo'):
            file_error('file_typo')
        codes = {value for name, value in vars(intake_module).items() if name.startswith('FILE_') and isinstance(value, str)}
        self.assertEqual(codes, set(FILE_MESSAGES))

    # Batches

    def test_batches_are_listed_for_the_company_only(self):
        self.commit(csv_bytes('Sara,0501234567'))
        self.env['liwan.crm.import.batch'].create({
            'company_id': self.other_company.id, 'user_id': self.manager.id, 'file_name': 'other.csv',
        })
        body = self.as_manager().api('GET', IMPORTS_URL).json()
        self.assertEqual(body['total'], 1)
        self.assertEqual([item['file_name'] for item in body['items']], ['customers.csv'])

    # Mapping profiles

    def test_mapping_profile_crud_and_reuse(self):
        manager = self.as_manager()
        created = manager.api('POST', MAPPINGS_URL, json={'name': 'Standard sheet', 'columns': {'name': 'Name', 'phone': 'Phone'}})
        self.assertEqual(created.status_code, 200, created.text)
        profile = created.json()
        self.assertEqual(profile['columns'], {'name': 'Name', 'phone': 'Phone', 'email': None, 'notes': None, 'package': None})
        self.assertEqual(set(profile), {'id', 'name', 'columns'})
        listed = manager.api('GET', MAPPINGS_URL).json()
        self.assertEqual((listed['total'], listed['items'][0]['id']), (1, profile['id']))
        body = {**self.body(csv_bytes('Sara,0501234567')), 'mapping_profile': profile['id']}
        del body['mapping']
        self.assertEqual(manager.api('POST', IMPORTS_URL, json=body).status_code, 200)
        self.assertEqual(manager.api('DELETE', f"{MAPPINGS_URL}/{profile['id']}").status_code, 200)
        self.assertEqual(manager.api('GET', MAPPINGS_URL).json()['total'], 0)

    def test_duplicate_profile_name_conflicts_per_company(self):
        manager = self.as_manager()
        payload = {'name': 'Standard sheet', 'columns': {'name': 'Name', 'phone': 'Phone'}}
        self.assertEqual(manager.api('POST', MAPPINGS_URL, json=payload).status_code, 200)
        again = manager.api('POST', MAPPINGS_URL, json=payload)
        self.assertEqual((again.status_code, again.json()['code']), (409, 'conflict'))
        self.env['liwan.crm.import.mapping'].create({
            'company_id': self.other_company.id, 'name': 'Other', 'name_column': 'A', 'phone_column': 'B',
        })
        self.assertEqual(manager.api('GET', MAPPINGS_URL).json()['total'], 1)

    def test_another_companys_profile_is_not_found(self):
        foreign = self.env['liwan.crm.import.mapping'].create({
            'company_id': self.other_company.id, 'name': 'Other', 'name_column': 'A', 'phone_column': 'B',
        })
        manager = self.as_manager()
        self.assertEqual(manager.api('DELETE', f'{MAPPINGS_URL}/{foreign.id}').status_code, 404)
        body = {**self.body(csv_bytes('Sara,0501234567')), 'mapping_profile': foreign.id}
        del body['mapping']
        self.assertEqual(manager.api('POST', IMPORTS_URL, json=body).status_code, 404)

    # Access

    def test_agent_without_import_capability_is_forbidden(self):
        content = csv_bytes('Sara,0501234567')
        agent = self.as_agent()
        self.assertEqual(self.preview(content, user=self.agent).status_code, 403)
        self.assertEqual(self.commit(content, user=self.agent).status_code, 403)
        for method, path in (('GET', IMPORTS_URL), ('GET', MAPPINGS_URL), ('POST', MAPPINGS_URL)):
            self.assertEqual(agent.api(method, path, json={}).status_code, 403, path)
        self.assertEqual(agent.api('DELETE', f'{MAPPINGS_URL}/1').status_code, 403)
        self.assertFalse(self.lead_of('+966501234567'))

    def test_closed_import_feature_refuses_every_route(self):
        set_probe_feature(self, key='crm.import', enabled=False)
        manager = self.as_manager()
        content = csv_bytes('Sara,0501234567')
        for method, path, json in (
            ('POST', PREVIEW_URL, self.body(content)), ('POST', IMPORTS_URL, self.body(content)), ('GET', IMPORTS_URL, None),
            ('GET', MAPPINGS_URL, None), ('POST', MAPPINGS_URL, {'name': 'x', 'columns': {'name': 'a', 'phone': 'b'}}),
        ):
            self.assertEqual(manager.api(method, path, json=json).status_code, 403, path)
        self.assertFalse(self.lead_of('+966501234567'))
